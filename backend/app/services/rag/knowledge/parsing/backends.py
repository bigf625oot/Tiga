import io
import os
import tempfile
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Protocol, Sequence

from fastapi import HTTPException
from PIL import Image

from app.core.config import settings

from .text import is_text_valid, sanitize_text


@dataclass(frozen=True)
class ParseChunk:
    page: int
    text: str
    meta: Optional[Dict[str, Any]] = None

    def to_dict(self) -> Dict[str, Any]:
        out: Dict[str, Any] = {"page": self.page, "text": self.text}
        if self.meta:
            out["meta"] = self.meta
        return out


class ParseBackend(Protocol):
    name: str

    def parse_path(self, file_path: str) -> List[ParseChunk]: ...

    def parse_bytes(
        self, content: bytes, filename: str, content_type: Optional[str]
    ) -> List[ParseChunk]: ...


def _pdf_chunks_from_text(text: str, backend: str) -> List[ParseChunk]:
    cleaned = sanitize_text(text or "").strip()
    if not cleaned:
        return []
    return [ParseChunk(page=1, text=cleaned, meta={"backend": backend})]


class DoclingPdfBackend:
    name = "docling"

    def _convert_to_text(self, converted: Any) -> str:
        doc = getattr(converted, "document", converted)
        for attr in ("export_to_text", "export_to_markdown"):
            fn = getattr(doc, attr, None)
            if callable(fn):
                return fn() or ""
        for attr in ("text", "markdown"):
            v = getattr(doc, attr, None)
            if isinstance(v, str):
                return v
        return str(doc) if doc is not None else ""

    def parse_path(self, file_path: str) -> List[ParseChunk]:
        try:
            from docling.document_converter import DocumentConverter
        except Exception as e:
            raise ImportError("docling not available") from e

        converter = DocumentConverter()
        converted = converter.convert(file_path)
        text = self._convert_to_text(converted)
        if not is_text_valid(text):
            return []
        return _pdf_chunks_from_text(text, self.name)

    def parse_bytes(
        self, content: bytes, filename: str, content_type: Optional[str]
    ) -> List[ParseChunk]:
        suffix = os.path.splitext(filename or "")[1] or ".pdf"
        with tempfile.NamedTemporaryFile(delete=True, suffix=suffix) as tmp:
            tmp.write(content)
            tmp.flush()
            return self.parse_path(tmp.name)


class PyMuPdfBackend:
    name = "pymupdf"

    def parse_path(self, file_path: str) -> List[ParseChunk]:
        try:
            import fitz
        except Exception as e:
            raise ImportError("fitz not available") from e

        doc = fitz.open(file_path)
        out: List[ParseChunk] = []
        try:
            for i in range(doc.page_count):
                p = doc.load_page(i)
                text = p.get_text("text") or ""
                text = sanitize_text(text).strip()
                if text:
                    out.append(ParseChunk(page=i + 1, text=text, meta={"backend": self.name}))
        finally:
            doc.close()

        if out and not is_text_valid("\n".join([c.text for c in out])):
            return []
        return out

    def parse_bytes(
        self, content: bytes, filename: str, content_type: Optional[str]
    ) -> List[ParseChunk]:
        try:
            import fitz
        except Exception as e:
            raise ImportError("fitz not available") from e

        doc = fitz.open(stream=content, filetype="pdf")
        out: List[ParseChunk] = []
        try:
            for i in range(doc.page_count):
                p = doc.load_page(i)
                text = p.get_text("text") or ""
                text = sanitize_text(text).strip()
                if text:
                    out.append(ParseChunk(page=i + 1, text=text, meta={"backend": self.name}))
        finally:
            doc.close()

        if out and not is_text_valid("\n".join([c.text for c in out])):
            return []
        return out


class PdfPlumberBackend:
    name = "pdfplumber"

    def parse_path(self, file_path: str) -> List[ParseChunk]:
        try:
            import pdfplumber
        except Exception as e:
            raise ImportError("pdfplumber not available") from e

        out: List[ParseChunk] = []
        with pdfplumber.open(file_path) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                text = sanitize_text(text).strip()
                if text:
                    out.append(ParseChunk(page=i + 1, text=text, meta={"backend": self.name}))

        if out and not is_text_valid("\n".join([c.text for c in out])):
            return []
        return out

    def parse_bytes(
        self, content: bytes, filename: str, content_type: Optional[str]
    ) -> List[ParseChunk]:
        try:
            import pdfplumber
        except Exception as e:
            raise ImportError("pdfplumber not available") from e

        out: List[ParseChunk] = []
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                text = sanitize_text(text).strip()
                if text:
                    out.append(ParseChunk(page=i + 1, text=text, meta={"backend": self.name}))

        if out and not is_text_valid("\n".join([c.text for c in out])):
            return []
        return out


class PyPdfBackend:
    name = "pypdf"

    def parse_path(self, file_path: str) -> List[ParseChunk]:
        try:
            import pypdf
        except Exception as e:
            raise ImportError("pypdf not available") from e

        out: List[ParseChunk] = []
        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for i, page in enumerate(reader.pages):
                try:
                    text = page.extract_text() or ""
                except Exception:
                    text = ""
                text = sanitize_text(text).strip()
                if text:
                    out.append(ParseChunk(page=i + 1, text=text, meta={"backend": self.name}))

        if out and not is_text_valid("\n".join([c.text for c in out])):
            return []
        return out

    def parse_bytes(
        self, content: bytes, filename: str, content_type: Optional[str]
    ) -> List[ParseChunk]:
        try:
            import pypdf
        except Exception as e:
            raise ImportError("pypdf not available") from e

        reader = pypdf.PdfReader(io.BytesIO(content))
        out: List[ParseChunk] = []
        for i, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
            text = sanitize_text(text).strip()
            if text:
                out.append(ParseChunk(page=i + 1, text=text, meta={"backend": self.name}))

        if out and not is_text_valid("\n".join([c.text for c in out])):
            return []
        return out


class OcrPdfBackend:
    name = "ocr"

    def parse_path(self, file_path: str) -> List[ParseChunk]:
        if not getattr(settings, "OCR_ENABLED", False):
            return []
        try:
            import fitz
            import pytesseract
        except Exception as e:
            raise ImportError("ocr dependencies not available") from e

        doc = fitz.open(file_path)
        out: List[ParseChunk] = []
        try:
            for i in range(doc.page_count):
                p = doc.load_page(i)
                pm = p.get_pixmap()
                img_bytes = pm.tobytes("png")
                img = Image.open(io.BytesIO(img_bytes))
                text = pytesseract.image_to_string(img) or ""
                text = sanitize_text(text).strip()
                if text:
                    out.append(ParseChunk(page=i + 1, text=text, meta={"backend": self.name, "ocr": True}))
        finally:
            doc.close()

        if out and not is_text_valid("\n".join([c.text for c in out])):
            return []
        return out

    def parse_bytes(
        self, content: bytes, filename: str, content_type: Optional[str]
    ) -> List[ParseChunk]:
        if not getattr(settings, "OCR_ENABLED", False):
            return []
        try:
            import fitz
            import pytesseract
        except Exception as e:
            raise ImportError("ocr dependencies not available") from e

        doc = fitz.open(stream=content, filetype="pdf")
        out: List[ParseChunk] = []
        try:
            for i in range(doc.page_count):
                p = doc.load_page(i)
                pm = p.get_pixmap()
                img_bytes = pm.tobytes("png")
                img = Image.open(io.BytesIO(img_bytes))
                text = pytesseract.image_to_string(img) or ""
                text = sanitize_text(text).strip()
                if text:
                    out.append(ParseChunk(page=i + 1, text=text, meta={"backend": self.name, "ocr": True}))
        finally:
            doc.close()

        if out and not is_text_valid("\n".join([c.text for c in out])):
            return []
        return out


class DocxBackend:
    name = "docx"

    def parse_path(self, file_path: str) -> List[ParseChunk]:
        try:
            import docx
        except Exception as e:
            raise ImportError("python-docx not available") from e

        doc = docx.Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs])
        text = sanitize_text(text).strip()
        if not text:
            return []
        return [ParseChunk(page=1, text=text, meta={"backend": self.name})]

    def parse_bytes(
        self, content: bytes, filename: str, content_type: Optional[str]
    ) -> List[ParseChunk]:
        try:
            import docx
        except Exception as e:
            raise ImportError("python-docx not available") from e

        doc = docx.Document(io.BytesIO(content))
        text = "\n".join([p.text for p in doc.paragraphs])
        text = sanitize_text(text).strip()
        if not text:
            return []
        return [ParseChunk(page=1, text=text, meta={"backend": self.name})]


class PlainTextBackend:
    name = "text"

    def parse_path(self, file_path: str) -> List[ParseChunk]:
        text = ""
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
        except Exception:
            try:
                with open(file_path, "r", encoding="gbk") as f:
                    text = f.read()
            except Exception:
                text = ""
        text = sanitize_text(text).strip()
        if not text:
            return []
        return [ParseChunk(page=1, text=text, meta={"backend": self.name})]

    def parse_bytes(
        self, content: bytes, filename: str, content_type: Optional[str]
    ) -> List[ParseChunk]:
        try:
            text = content.decode("utf-8")
        except Exception:
            try:
                text = content.decode("gbk")
            except Exception as e:
                raise HTTPException(
                    status_code=400,
                    detail="无法解码文件内容，请确保是 UTF-8 或 GBK 编码的文本文件。",
                ) from e
        text = sanitize_text(text).strip()
        if not text:
            return []
        return [ParseChunk(page=1, text=text, meta={"backend": self.name})]


def build_default_backends(order: Sequence[str]) -> List[ParseBackend]:
    mapping: Dict[str, ParseBackend] = {
        "docling": DoclingPdfBackend(),
        "pymupdf": PyMuPdfBackend(),
        "pdfplumber": PdfPlumberBackend(),
        "pypdf": PyPdfBackend(),
        "ocr": OcrPdfBackend(),
        "docx": DocxBackend(),
        "text": PlainTextBackend(),
    }
    out: List[ParseBackend] = []
    for name in order:
        key = (name or "").strip().lower()
        b = mapping.get(key)
        if b:
            out.append(b)
    return out
